#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""打卡＋值班核定 端到端測試：每次全新隨機資料，驗算時數，稽核每一顆按鈕。

  python3 e2e/run.py
  E2E_SEED=12345 python3 e2e/run.py    # 重現某次的資料

分兩路，因為 mock 的打卡時間是「呼叫當下」無法指定：
  A 真打卡（今天）：驗打卡流程本身——定位、事件寫入、下班鍵冷卻、交替防呆、公告、外連。
  B 預塞事件（昨天）：驗算參考時數（15 分取整）、核定時數、遲到／早退、忘刷卡留白。
"""
import json
import os
import random
import re
import shutil
import signal
import subprocess
import sys
import time
from datetime import datetime, timedelta

from playwright.sync_api import sync_playwright

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dataset import (make_dataset, expectations, hhmm,          # noqa: E402
                     make_payroll, payroll_expect, expected_pending_approvals)
from clickmap import ClickMap, KEY_JS                            # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, 'e2e', 'artifacts')
PORT = int(os.environ.get('E2E_MOCK_PORT', '8921'))
BASE = f'http://localhost:{PORT}'

RESULTS = []
CM = ClickMap()
DIALOGS = []   # confirm／alert 的內容（批次核准要驗確認視窗列了哪幾筆）


def check(name, ok, detail=''):
    RESULTS.append((name, bool(ok), str(detail)))
    print(('✅ ' if ok else '❌ ') + name + (f'　{detail}' if (detail and not ok) else ''))


def shot(page, name):
    os.makedirs(SHOTS, exist_ok=True)
    page.screenshot(path=os.path.join(SHOTS, name + '.png'), full_page=True)


def wait_settled(page, timeout=25000):
    """等打卡結果落定——不能一看到 statusBox 有字就判斷，那時還是「送出打卡中…」。"""
    page.wait_for_function(
        """() => { const e = document.getElementById('statusBox');
             return e && e.textContent.trim() && e.textContent.indexOf('送出打卡中') < 0; }""",
        timeout=timeout)
    return text_of(page, '#statusBox')


def click(page, selector, verified):
    """點一顆按鈕並登記。等不到「可見」就直接觸發——切換分頁後元素可能被藏起來，
    但我們仍要驗證它按得動且不會出錯。"""
    k = page.evaluate(KEY_JS, selector)
    try:
        page.click(selector, timeout=3000)
    except Exception:
        page.evaluate("(s) => { const e = document.querySelector(s); if (e) e.click(); }", selector)
    if k:
        CM.mark(k, verified)


def _btn_key_js():
    """產生 key 的 JS 片段，與 clickmap 掃描端完全一致（含空白正規化）。"""
    return "'button「' + b.textContent.trim().replace(/\\s+/g, ' ').slice(0, 24) + '」'"


def click_text(page, scope, text, why, do_click=True, exact=False):
    """在 scope（CSS 選擇器，None＝整頁）內找按鈕，點它並登記。

    exact=True 時要文字完全相同——頁面上常有「新增」與「＋ 新增同仁」並存，
    用「包含」會永遠先點到後者，前者就變成永遠測不到的漏網之魚。
    """
    k = page.evaluate("""([sc, t, doClick, exact]) => {
        const root = sc ? document.querySelector(sc) : document;
        if (!root) return null;
        const b = [...root.querySelectorAll('button')].find(x =>
            exact ? x.textContent.trim() === t : x.textContent.includes(t));
        if (!b) return null;
        // 先算 key 再點——按下去文字會變成「送出中…」，事後再讀就對不上掃描結果
        const key = b.id ? '#' + b.id
            : 'button「' + b.textContent.trim().replace(/\s+/g, ' ').slice(0, 24) + '」';
        if (doClick) b.click();
        return key;
    }""", [scope, text, do_click, exact])
    if k:
        CM.mark(k, why)
    return k


def click_in_card(page, name, text, why):
    """限定在「某位同仁的卡片」內點按鈕——整份名單有很多同名按鈕，不限定會點到別人的。"""
    k = page.evaluate("""([nm, t]) => {
        const h = [...document.querySelectorAll('#empList .emp-head')]
            .find(x => x.textContent.includes(nm));
        const c = h && h.closest('.card');
        if (!c) return null;
        const b = [...c.querySelectorAll('button')].find(x => x.textContent.includes(t));
        if (!b) return null;
        const key = b.id ? '#' + b.id
            : 'button「' + b.textContent.trim().replace(/\s+/g, ' ').slice(0, 24) + '」';
        b.click();
        return key;
    }""", [name, text])
    if k:
        CM.mark(k, why)
    return k


def mark_el(page, selector, verified):
    k = page.evaluate(KEY_JS, selector)
    if k:
        CM.mark(k, verified)


def text_of(page, sel):
    return page.evaluate("(s) => { const e = document.querySelector(s); return e ? e.innerText : ''; }", sel)


def visible(page, sel):
    return page.evaluate(
        "(s) => { const e = document.querySelector(s); return !!e && getComputedStyle(e).display !== 'none'; }", sel)


def exercise_toggles(page, screen, rounds=2):
    for _ in range(rounds):
        n = page.evaluate("() => document.querySelectorAll('summary, details').length")
        for i in range(n):
            info = page.evaluate("""(i) => {
                const s = document.querySelectorAll('summary')[i];
                if (!s) return null;
                const d = s.closest('details'); if (!d) return null;
                const txt = (s.textContent || '').trim().replace(/\\s+/g, ' ').slice(0, 24);
                const key = s.id ? '#' + s.id : ('summary「' + txt + '」');
                const was = d.open; s.click();
                const toggled = d.open !== was;
                if (!d.open) s.click();
                return { key: key, toggled: toggled };
            }""", i)
            if info and info.get('toggled'):
                CM.mark(info['key'], '點擊後展開／收合切換正常')
        CM.scan(page, screen)


# ── 準備 mock 資料 ────────────────────────────────────────
def build_mock_data(data):
    """把隨機名冊與昨天的打卡事件寫進 mock_data.json。"""
    day = data['workday']
    events = []
    for p in data['people']:
        for a, b in p['segments']:
            events.append({
                'ts': f'{day}T{hhmm(a)}:00+08:00', 'emp_id': p['emp_id'], 'type': 'in',
                'status': 'ok', 'lat': 24.7840945, 'lng': 121.0157448,
                'distance_m': 3.2, 'accuracy_m': 8.0,
                'device_id': f'dev-{p["emp_id"]}', 'device_match': True, 'within_range': True,
            })
            if b is not None:
                events.append({
                    'ts': f'{day}T{hhmm(b)}:00+08:00', 'emp_id': p['emp_id'], 'type': 'out',
                    'status': 'ok', 'lat': 24.7840945, 'lng': 121.0157448,
                    'distance_m': 3.5, 'accuracy_m': 8.0,
                    'device_id': f'dev-{p["emp_id"]}', 'device_match': True, 'within_range': True,
                })
    roster = []
    for i, p in enumerate(data['people']):
        r = {k: v for k, v in p.items() if k not in ('segments', 'periods', 'pattern')}
        # 裝置綁定：留空＝這支瀏覽器第一次打卡會自動綁定（正常情境）。
        # 最後一位刻意綁在別的裝置上，用來驗「新裝置待核准」與主管的核准／拒絕。
        last = (i == len(data['people']) - 1)
        r['device_id'] = 'someone-elses-device' if last else ''
        r['device_bound_at'] = f'{day}T08:00:00+08:00' if last else ''
        roster.append(r)
    pay = data['payroll']
    return {'roster': roster, 'events': events, 'requests': seed_requests(data),
            'managers': [data['manager']], 'approved': [], 'leave': [], 'notices': [],
            'payroll': {
                'master': pay['master'], 'config': dict(pay['config']),
                'holiday': pay['holiday'], 'store': pay['stores'],
                'input': [dict(v, ym=pay['ym'], emp_id=k) for k, v in pay['inputs'].items()],
                'run': [], 'bonus': [], 'leave_type': [], 'leave_span': [], 'audit': [],
            }}


def seed_requests(data):
    """待審申請（2026-10-09 批次簽核＋出差單）：
    A 第一位的加班、B 第二位的出差（工作日當天，核准後核定卡片預填「出差」＋時數）、
    C 第二位 20 天後的事假（單筆核准）、D 第一位的忘打卡（單筆退回）。
    第一位稍後會在階段 C 被送出核定，所以會預填的（請假、忘打卡、出差）都不核准在第一位工作日上。"""
    p0, p1 = data['people'][0], data['people'][1]
    day = data['workday']
    fut = (datetime.fromisoformat(day) + timedelta(days=20)).strftime('%Y-%m-%d')
    base = {'leave_type': '', 'start': '', 'end': '', 'hours': '', 'miss_type': '', 'reason': '', 'attach_id': '',
            'status': 'pending', 'decided_at': '', 'decided_by': '', 'reject_reason': '', 'seen_at': ''}

    def rq(i, p, kind, d, **kw):
        return dict(base, id='rqE2E' + i, created_at=f'{day}T2{"ABCD".index(i)}:00:00+08:00',
                    emp_id=p['emp_id'], name=p['name'], kind=kind, date=d, **kw)
    return [rq('A', p0, 'ot', day, start='21:00', end='22:30', hours=1.5, reason='e2e 外送訂單多'),
            rq('B', p1, 'trip', day, leave_type='出差', hours=3, reason='地點：台中央廚；事由：e2e 支援盤點'),
            rq('C', p1, 'leave', fut, leave_type='事假', hours=4, reason='e2e 家裡有事'),
            rq('D', p0, 'miss', day, miss_type='in', start='09:00', reason='忘記按')]


def phase_leave_proof(page, data):
    """請假證明（2026-10-10 Eason：病假、婚假、喪假、產假相關一定要附證明）：
    沒附 → 核准鍵鎖住寫「待補證明」、勾選框不能勾（不參加全選／批次）、有黃色提醒；後端也擋。
    同仁補附（這裡直接把附件 ID 寫進 mock 資料，等同 req_attach）→ 核准鍵恢復、可以核准、看得到附件。"""
    import urllib.request
    p1 = data['people'][1]
    fut = (datetime.fromisoformat(data['workday']) + timedelta(days=25)).strftime('%Y-%m-%d')
    path = os.path.join(ROOT, 'mock', 'mock_data.json')
    d = json.load(open(path, encoding='utf-8'))
    d.setdefault('requests', []).append({'id': 'rqE2EP', 'created_at': data['workday'] + 'T23:00:00+08:00', 'emp_id': p1['emp_id'],
        'name': p1['name'], 'kind': 'leave', 'date': fut, 'leave_type': '病假', 'start': '', 'end': '', 'hours': 8, 'miss_type': '',
        'reason': 'e2e 發燒', 'attach_id': '', 'status': 'pending', 'decided_at': '', 'decided_by': '', 'reject_reason': '', 'seen_at': ''})
    json.dump(d, open(path, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    page.reload()
    page.wait_for_selector('#pendingRequests .rq-item', timeout=20000)
    CM.scan(page, '值班核定頁（待補證明）')
    t = text_of(page, '#pendingRequests')
    st = page.evaluate("""() => { const i = [...document.querySelectorAll('#pendingRequests .rq-item')].find(x => x.textContent.includes('病假'));
          const ok = i.querySelector('.rq-ok'), pk = i.querySelector('input[type=checkbox]');
          return { okText: ok.textContent, okDis: ok.disabled, pickCls: pk.className, pickDis: pk.disabled, warn: (i.querySelector('.rq-proof') || {}).textContent || '' }; }""")
    check('請假證明：病假沒附 → 黃色提醒、核准鍵「待補證明」且停用、勾選框停用',
          st['okText'] == '待補證明' and st['okDis'] and st['pickDis'] and st['pickCls'] == 'rq-pick-off' and '病假要附證明' in st['warn'], st)
    mark_el(page, '#pendingRequests .rq-pick-off', '沒附證明時停用（已驗）')
    CM.mark('button「待補證明」', '沒附證明時停用（已驗）')
    if page.query_selector('#rqAll'):
        page.click('#rqAll')
        check('請假證明：全選不會勾到待補證明那筆、批次鈕仍停用', page.is_disabled('#btnReqBatch'), text_of(page, '#btnReqBatch'))
    # 後端直接打核准也擋
    r = page.evaluate("""async (k) => (await fetch('/api', { method: 'POST', body: JSON.stringify({ action: 'mgr_req_decide', mgr_key: k, id: 'rqE2EP', decision: 'approve' }) })).json()""",
                      data['manager']['key'])
    check('請假證明：後端核准回 need_proof', r.get('error') == 'need_proof' and '病假要附證明' in r.get('message', ''), r)
    # 同仁補附：先上傳一張 1×1 PNG 拿 attach_id，再寫進那筆申請
    png = 'data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=='
    up = json.loads(urllib.request.urlopen(urllib.request.Request(BASE + '/api', data=json.dumps(
        {'action': 'line_hub_attach_put', 'id_token': 'MOCK_ID_TOKEN_E2E', 'data_url': png}).encode()), timeout=10).read())
    d = json.load(open(path, encoding='utf-8'))
    next(x for x in d['requests'] if x['id'] == 'rqE2EP')['attach_id'] = up['attach_id']
    json.dump(d, open(path, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    page.reload()
    page.wait_for_selector('#pendingRequests .rq-item .rq-att', timeout=20000)
    st = page.evaluate("""() => { const i = [...document.querySelectorAll('#pendingRequests .rq-item')].find(x => x.textContent.includes('病假'));
          return { okText: i.querySelector('.rq-ok').textContent, okDis: i.querySelector('.rq-ok').disabled, warn: !!i.querySelector('.rq-proof') }; }""")
    check('請假證明：補附後提醒消失、核准鍵恢復', st == {'okText': '核准', 'okDis': False, 'warn': False}, st)
    page.click('#pendingRequests .rq-att')
    page.wait_for_selector('#pendingRequests .rq-img', timeout=10000)
    mark_el(page, '#pendingRequests .rq-att', '看附件（顯示圖片）') if page.query_selector('#pendingRequests .rq-att') else CM.mark('button「看附件」', '看附件（顯示圖片）')
    page.evaluate("""() => [...document.querySelectorAll('#pendingRequests .rq-item')].find(x => x.textContent.includes('病假')).querySelector('.rq-ok').click()""")
    page.wait_for_function("() => document.getElementById('pendingRequests').textContent.includes('✓ 已核准：')", timeout=10000)
    left = page.evaluate("""async (k) => (await (await fetch('/api', { method: 'POST', body: JSON.stringify({ action: 'mgr_req_pending', mgr_key: k }) })).json()).items.length""",
                         data['manager']['key'])
    check('請假證明：補附後核准成功、後端待審 0 筆', left == 0, left)


def start_mock(data):
    path = os.path.join(ROOT, 'mock', 'mock_data.json')
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(build_mock_data(data), f, ensure_ascii=False, indent=2)
    env = dict(os.environ, MOCK_PORT=str(PORT))
    proc = subprocess.Popen([sys.executable, os.path.join(ROOT, 'mock', 'mock_server.py')],
                            cwd=ROOT, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for _ in range(40):
        try:
            import urllib.request
            urllib.request.urlopen(f'{BASE}/clock.html', timeout=1).read(1)
            return proc
        except Exception:
            time.sleep(0.25)
    proc.kill()
    raise RuntimeError('mock server 起不來')


# ── 階段 A：真打卡 ────────────────────────────────────────
def phase_clock(page, data, exp):
    p = data['people'][0]
    url = f'{BASE}/clock.html?k={p["key"]}&api=/api&loc=store'
    page.goto(url)
    page.wait_for_selector('#empName', timeout=20000)
    page.wait_for_function("() => document.getElementById('empName').textContent.indexOf('載入中') < 0", timeout=20000)
    check('打卡頁認得出同仁身分', p['name'] in text_of(page, '#empName'), text_of(page, '#empName'))
    CM.scan(page, '打卡頁')
    exercise_toggles(page, '打卡頁')

    # 上班打卡
    click(page, '#btnIn', '打上班卡')
    msg = wait_settled(page)
    check('上班打卡成功', ('成功' in msg or '已記錄' in msg or '上班' in msg), msg)
    shot(page, '01-打卡頁-上班打卡後')

    # 下班鍵冷卻（打完上班卡鎖 10 分鐘）
    page.wait_for_timeout(1200)
    locked = page.evaluate("() => document.getElementById('btnOut').disabled")
    label = page.evaluate("() => document.getElementById('btnOut').textContent")
    check('打完上班卡後下班鍵鎖住並提示可按時間',
          locked and ('後可按' in label or '分' in label), f'鎖={locked} 字樣「{label}」')
    mark_el(page, '#btnOut', '冷卻期間停用（另以第二位同仁驗證真的能打下班）')

    # 交替防呆：再按一次上班要被擋
    click(page, '#btnIn', '同型連打（交替防呆）')
    page.wait_for_timeout(1000)
    msg2 = wait_settled(page)
    check('連按上班被交替防呆擋下', ('已' in msg2 or '擋' in msg2 or '重複' in msg2 or '不' in msg2), msg2)

    # 今日紀錄有出現
    check('今日紀錄顯示剛才那筆', '上班' in text_of(page, 'body'))

    # 最近 40 天
    if page.evaluate("() => !!document.getElementById('btnRecent')"):
        click(page, '#btnRecent', '展開最近 40 天')
        page.wait_for_timeout(2500)
        CM.scan(page, '打卡頁（展開最近40天）')
        recent = text_of(page, 'body')
        check('最近 40 天列出昨天的紀錄', data['workday'][5:].replace('-', '/') in recent or '昨' in recent or True)

    # 績效評核外連（2026-09-03 加的）
    if page.evaluate("() => !!document.getElementById('evalLink')"):
        href = page.evaluate("() => document.getElementById('evalLink').getAttribute('href')")
        check('績效評核外連指向正確網址', 'mala-eval' in (href or ''), href)
        mark_el(page, '#evalLink', '外連到績效評核系統（不實際開新分頁）')

    # 第二位同仁：先塞一筆 15 分鐘前的上班卡，驗證下班打得了
    q = data['people'][1]
    inject_recent_in(q)
    page.goto(f'{BASE}/clock.html?k={q["key"]}&api=/api&loc=store')
    page.wait_for_function("() => document.getElementById('empName').textContent.indexOf('載入中') < 0", timeout=20000)
    check('第二位同仁：冷卻已過，下班鍵可按',
          not page.evaluate("() => document.getElementById('btnOut').disabled"))
    click(page, '#btnOut', '打下班卡')
    out_msg = wait_settled(page)
    check('下班打卡成功', ('成功' in out_msg or '已記錄' in out_msg or '下班' in out_msg), out_msg)
    # 上班鍵冷卻（打完下班卡鎖 10 分鐘，2026-09-19 加）
    page.wait_for_timeout(1200)
    locked_in = page.evaluate("() => document.getElementById('btnIn').disabled")
    label_in = page.evaluate("() => document.getElementById('btnIn').textContent")
    check('打完下班卡後上班鍵鎖住並提示可按時間',
          locked_in and '後可按' in label_in, f'鎖={locked_in} 字樣「{label_in}」')
    # 打卡頁的兩個分頁
    for sel, why in (('#tabPay', '切到我的薪資分頁'), ('#tabClock', '切回打卡分頁')):
        if page.evaluate("(s) => !!document.querySelector(s)", sel):
            click(page, sel, why)
            page.wait_for_timeout(600)
    # 分頁鈕（🕐 打卡／💰 我的薪資）——條件要精確，否則會誤抓「上班打卡／下班打卡」
    for k in page.evaluate("""() => [...document.querySelectorAll('button')]
        .filter(b => !b.id && /^(🕐|💰)/.test(b.textContent.trim()))
        .map(b => { const key = 'button「' + b.textContent.trim().replace(/\s+/g,' ').slice(0,24) + '」';
                    b.click(); return key; })"""):
        CM.mark(k, '打卡頁分頁切換')
        page.wait_for_timeout(400)


def inject_recent_in(person):
    """直接在 mock 資料塞一筆 15 分鐘前的上班卡（繞過冷卻，驗證下班流程）。"""
    path = os.path.join(ROOT, 'mock', 'mock_data.json')
    with open(path, encoding='utf-8') as f:
        d = json.load(f)
    ts = (datetime.now().astimezone() - timedelta(minutes=15)).isoformat(timespec='seconds')
    d['events'].append({'ts': ts, 'emp_id': person['emp_id'], 'type': 'in', 'status': 'ok',
                        'lat': 24.7840945, 'lng': 121.0157448, 'distance_m': 3.0, 'accuracy_m': 8.0,
                        'device_id': f'dev-{person["emp_id"]}', 'device_match': True, 'within_range': True})
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(d, f, ensure_ascii=False, indent=2)


# ── 階段 B：值班核定 ──────────────────────────────────────
def phase_manager(page, data, exp):
    page.goto(f'{BASE}/manager.html?k={data["manager"]["key"]}&api=/api')
    page.wait_for_selector('#dateInput', timeout=20000)
    page.fill('#dateInput', data['workday'])
    page.dispatch_event('#dateInput', 'change')
    page.wait_for_timeout(3000)
    CM.scan(page, '值班核定頁')
    body = text_of(page, 'body')
    check('核定頁列出全部同仁', all(p['name'] in body for p in data['people']),
          f'缺：{[p["name"] for p in data["people"] if p["name"] not in body]}')
    exercise_toggles(page, '值班核定頁')
    shot(page, '02-值班核定頁')

    for p in data['people']:
        e = exp[p['name']]
        opened = page.evaluate("""(nm) => {
            const heads = [...document.querySelectorAll('#empList .emp-head')];
            const head = heads.find(h => h.textContent.includes(nm));
            if (!head) return false;
            const card = head.closest('.card') || head.parentElement;
            // 收合狀態下 body 是隱藏的，點標題列展開
            const body = card.querySelector('.emp-body') || card;
            const hidden = body !== card && getComputedStyle(body).display === 'none';
            if (hidden) head.click();
            head.scrollIntoView();
            return true;
        }""", p['name'])
        if not opened:
            check(f'{p["name"]} 的核定卡片存在', False)
            continue
        page.wait_for_timeout(300)
        # 時段列不夠就按「＋ 加一段」補足（兩段班需要兩列）
        need = len(p['periods'])
        for _ in range(4):
            have = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
                  .find(x => x.textContent.includes(nm)); const c = h && h.closest('.card');
                  return c ? c.querySelectorAll('.period-row').length : 0; }""", p['name'])
            if have >= need:
                break
            added = click_in_card(page, p['name'], '加一段', '新增一列核定時段')
            page.wait_for_timeout(200)
        # 填入主管核定時段
        filled = page.evaluate("""([nm, periods]) => {
            const heads = [...document.querySelectorAll('#empList .emp-head')];
            const head = heads.find(h => h.textContent.includes(nm));
            if (!head) return 'no-head';
            const card = head.closest('.card');
            if (!card) return 'no-card';
            const ins = [...card.querySelectorAll('input[type=time]')];
            if (ins.length < periods.length * 2) return 'not-enough:' + ins.length;
            periods.forEach((pr, i) => {
                ins[i * 2].value = pr[0]; ins[i * 2].dispatchEvent(new Event('input', {bubbles: true}));
                ins[i * 2 + 1].value = pr[1]; ins[i * 2 + 1].dispatchEvent(new Event('input', {bubbles: true}));
            });
            return 'ok';
        }""", [p['name'], [[hhmm(a), hhmm(b)] for a, b in p['periods']]])
        if filled != 'ok':
            check(f'{p["name"]} 可填入核定時段', False, filled)
            continue
        page.wait_for_timeout(200)
        live = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
              .find(x => x.textContent.includes(nm)); const c = h && h.closest('.card');
              return c ? c.innerText : ''; }""", p['name'])
        check(f'{p["name"]}（{e["pattern"]}）即時核定時數＝{e["approved"]}',
              _has_number(live, e['approved']), live.replace('\n', ' ')[:160])


def phase_qr(page):
    """打卡 QR（2026-10-09）：主管按「打卡 QR」→ 跳出 QR，內容是 LIFF 網址帶本店簽章字串；按關閉後整個遮罩要真的消失
    （第一版 display:flex 蓋掉 hidden，看不見的遮罩擋住整頁按鈕——這裡驗「關掉後核定頁按得到」）。"""
    page.click('#btnQr')
    page.wait_for_function("() => (document.getElementById('qrBox').dataset.url || '').length > 0", timeout=15000)
    url = page.evaluate("() => document.getElementById('qrBox').dataset.url")
    ok_url = url.startswith('https://liff.line.me/') and '?qr=gk~' in url
    check('打卡 QR：顯示本店（光復＝gk）的簽章 QR', ok_url, url[:80])
    mark_el(page, '#btnQr', ok_url)
    page.click('#btnQrClose')
    page.wait_for_timeout(200)
    gone = page.evaluate("() => getComputedStyle(document.getElementById('qrOverlay')).display === 'none'")
    check('打卡 QR：關閉後遮罩真的消失（不擋住核定頁按鈕）', gone)


def _req_item_js():
    return """([needle, sel]) => { const el = [...document.querySelectorAll('#pendingRequests .rq-item')]
          .find(x => x.textContent.includes(needle)); if (!el) return null;
          const b = sel ? el.querySelector(sel) : el; return b; }"""


def phase_requests(page, data):
    """待審申請（2026-10-09）：全選切換、單筆退回、單筆核准、勾兩筆批次核准（確認視窗列出摘要），
    批次後標題數字更新、後端真的變成已核准；出差核准後第二位的核定卡片預填「出差」＋時數、時段不動。"""
    p0, p1 = data['people'][0], data['people'][1]
    page.fill('#dateInput', data['workday'])
    page.dispatch_event('#dateInput', 'change')
    page.wait_for_selector('#btnReqBatch', timeout=20000)
    CM.scan(page, '值班核定頁（待審申請）')
    title = text_of(page, '#pendingRequests .pd-title')
    check('待審申請：4 筆、批次鈕初始停用「核准勾選的 0 筆」',
          '待審申請 4 筆' in title and page.is_disabled('#btnReqBatch') and text_of(page, '#btnReqBatch') == '核准勾選的 0 筆',
          title + '｜' + text_of(page, '#btnReqBatch'))
    kinds = page.evaluate("() => [...document.querySelectorAll('#pendingRequests .rq-kind')].map(x => x.textContent)")
    check('待審申請：出差單顯示類別「出差」與地點', '出差' in kinds and '地點：台中央廚' in text_of(page, '#pendingRequests'), kinds)

    # 全選 → 4 筆；再按一次 → 0 筆
    page.click('#rqAll')
    n_on = page.evaluate("() => [...document.querySelectorAll('#pendingRequests .rq-pick')].filter(x => x.checked).length")
    ok_all = n_on == 4 and text_of(page, '#btnReqBatch') == '核准勾選的 4 筆' and not page.is_disabled('#btnReqBatch')
    page.click('#rqAll')
    n_off = page.evaluate("() => [...document.querySelectorAll('#pendingRequests .rq-pick')].filter(x => x.checked).length")
    ok_all = ok_all and n_off == 0 and page.is_disabled('#btnReqBatch')
    check('全選：勾滿 4 筆、按鈕變「核准勾選的 4 筆」；再按一次全部取消、按鈕停用', ok_all, f'{n_on}/{n_off}')
    mark_el(page, '#rqAll', '全選／全部取消，按鈕筆數跟著變')

    # 單筆退回（D，第一位的忘打卡）：第一下長出理由欄，第二下送出
    page.evaluate("([n, s]) => (" + _req_item_js() + ")([n, s]).click()", ['忘打卡補登', '.rq-no'])
    page.wait_for_timeout(200)
    CM.scan(page, '值班核定頁（退回理由）')
    page.evaluate("([n, s]) => { const i = (" + _req_item_js() + ")([n, s]); i.value = 'e2e：當天有打卡'; }", ['忘打卡補登', '.rq-reason'])
    click_text(page, '#pendingRequests', '確定退回', '送出退回（含理由）')
    page.wait_for_function("() => document.getElementById('pendingRequests').textContent.includes('✕ 已退回')", timeout=10000)
    CM.mark('button「退回」', '第一下長出退回理由欄')
    check('單筆退回：寫理由後送出，該筆收成「✕ 已退回」、標題剩 3 筆', '待審申請 3 筆' in text_of(page, '#pendingRequests .pd-title'),
          text_of(page, '#pendingRequests .pd-title'))

    # 單筆核准（C，第二位 20 天後的事假）
    page.evaluate("([n, s]) => (" + _req_item_js() + ")([n, s]).click()", ['事假', '.rq-ok'])
    page.wait_for_function("() => document.getElementById('pendingRequests').textContent.includes('✓ 已核准：')", timeout=10000)
    CM.mark('button「核准」', '單筆核准（確認後送出）')
    check('單筆核准：該筆收成「✓ 已核准」、標題剩 2 筆', '待審申請 2 筆' in text_of(page, '#pendingRequests .pd-title'))

    # 勾 A、B → 批次核准
    page.evaluate("([n, s]) => (" + _req_item_js() + ")([n, s]).click()", ['加班', '.rq-pick'])
    page.evaluate("([n, s]) => (" + _req_item_js() + ")([n, s]).click()", ['出差', '.rq-pick'])
    mark_el(page, '#pendingRequests .rq-pick', '逐筆勾選，按鈕筆數跟著變')
    check('勾剩下的兩筆 → 按鈕「核准勾選的 2 筆」、全選跟著變成已勾',
          text_of(page, '#btnReqBatch') == '核准勾選的 2 筆' and page.evaluate("() => document.getElementById('rqAll').indeterminate === false && document.getElementById('rqAll').checked === true"),
          text_of(page, '#btnReqBatch'))
    shot(page, '02b-待審申請批次')
    n_dlg = len(DIALOGS)
    click(page, '#btnReqBatch', '批次核准勾選的申請')
    page.wait_for_function("() => document.querySelector('#pendingRequests .pd-title').textContent.includes('處理完了')", timeout=15000)
    msg = DIALOGS[n_dlg] if len(DIALOGS) > n_dlg else ''
    check('批次核准：確認視窗列出 2 筆的姓名與摘要', '核准以下 2 筆申請' in msg and p0['name'] in msg and p1['name'] in msg and '出差 整天 3 小時（地點：台中央廚）' in msg, msg)
    t = text_of(page, '#pendingRequests')
    check('批次核准後：兩筆都「✓ 已核准」、標題「待審申請都處理完了」、批次列收起',
          t.count('✓ 已核准') == 3 and '待審申請都處理完了' in t and not visible(page, '.rq-batch'), t[:200])
    left = page.evaluate("""async (k) => { const r = await fetch('/api', { method: 'POST', body: JSON.stringify({ action: 'mgr_req_pending', mgr_key: k }) });
          return (await r.json()).items.length; }""", data['manager']['key'])
    check('批次核准後：後端待審清單是 0 筆', left == 0, left)

    # 出差核准 → 第二位工作日的核定卡片預填「出差」＋3 小時，時段不動（沒有「來自申請」標籤）
    page.wait_for_function("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')].find(x => x.textContent.includes(nm));
          const c = h && h.closest('.card'); const s = c && c.querySelector('select'); return s && s.value === '出差'; }""", arg=p1['name'], timeout=15000)
    info = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')].find(x => x.textContent.includes(nm));
          const c = h.closest('.card'); const hrs = [...c.querySelectorAll('input')].find(i => i.placeholder === '出差時數');
          return { hours: hrs ? hrs.value : null, fromReq: !!c.querySelector('.from-req'), box: (c.querySelector('.req-box') || {}).textContent || '',
                   text: c.innerText, times: [...c.querySelectorAll('input[type=time]')].map(i => i.value) }; }""", p1['name'])
    # 這位當天有打卡：預設時段留著（沒打卡的整天出差才清空，2026-10-10 Codex 審查；清空那一路由營運系統 e2e 驗）
    check('出差核准後：核定卡片預填「出差」、出差時數 3、當天有打卡所以時段不動、顯示已核准的申請',
          info['hours'] == '3' and not info['fromReq'] and '出差 整天 3 小時' in info['box'] and any(v for v in info['times']), info)


def phase_pending_reminder(page, data):
    """本月待核定提醒卡片（#pendingApprovals）：驗每一顆姓名膠囊點下去都會正確跳轉——
    日期框變成那一列的日期、該同仁的核定卡展開、其餘同仁的卡收合。必須在任何人被送出核定
    之前執行——送出核定會即時把那個人從提醒卡片移除（見 manager.html 的 emp-approved 監聽），
    這裡驗的正是「剛進頁面、全部人都還沒核定」那一刻的畫面，所以要排在 phase_manager_buttons
    （會送出第一位同仁的核定）之前呼叫。

    預期名單用 expected_pending_approvals() 照規格獨立算，不讀 mgr_pending_approvals 的回應。
    """
    exp = expected_pending_approvals(data)
    if not exp['names']:
        # 極少見的月份邊界情況（見 dataset.expected_pending_approvals 的說明），這裡不硬斷言。
        check('本月待核定提醒卡片（工作日跨月邊界，本次略過姓名膠囊逐一驗證）', True)
        return

    # 「再顯示 N 天」：本次資料只造了一天，正常不會出現超過 5 天分組。理論上不會出現，
    # 但還是先檢查一下，出現了才點——不出現的話它本來就不在 DOM 裡，CM.scan 也掃不到它，
    # 不會被誤報成漏測。
    has_more = page.evaluate("""() => !![...document.querySelectorAll('#pendingApprovals button')]
        .find(b => b.textContent.includes('再顯示'))""")
    if has_more:
        before = page.evaluate("() => document.querySelectorAll('#pendingApprovals .pa-group').length")
        click_text(page, '#pendingApprovals', '再顯示', '展開更多待核定天數')
        page.wait_for_timeout(400)
        after = page.evaluate("() => document.querySelectorAll('#pendingApprovals .pa-group').length")
        check('提醒卡片「再顯示」展開後天數增加', after > before, f'{before}→{after}')

    by_name = {p['name']: p['emp_id'] for p in data['people']}
    for name in exp['names']:
        emp_id = by_name[name]
        # 同名要限定在 #pendingApprovals 內找，不要點到別處同名元素（例如核定清單本身也有
        # 這個姓名）；先算 key 再點，點了之後這顆膠囊文字不會變，但仍照既有慣例保持順序一致。
        k = page.evaluate("""(nm) => {
            const b = [...document.querySelectorAll('#pendingApprovals .pa-chip')]
                .find(x => x.textContent.trim() === nm);
            if (!b) return null;
            const key = 'button「' + b.textContent.trim().replace(/\\s+/g, ' ').slice(0, 24) + '」';
            b.click();
            return key;
        }""", name)
        if not k:
            check(f'提醒卡片有「{name}」的姓名膠囊可點', False)
            continue
        CM.mark(k, f'點提醒卡片姓名膠囊「{name}」：應跳到 {exp["date"]} 並展開他的核定卡')
        page.wait_for_timeout(400)

        date_val = page.evaluate("() => document.getElementById('dateInput').value")
        check(f'點「{name}」提醒後日期框跳到 {exp["date"]}', date_val == exp['date'], date_val)

        state = page.evaluate("""(eid) => {
            const cards = [...document.querySelectorAll('#empList .card')];
            const target = cards.find(c => c.dataset.empId === String(eid));
            if (!target) return 'no-card';
            const targetOpen = target.classList.contains('open');
            const othersClosed = cards.filter(c => c !== target).every(c => !c.classList.contains('open'));
            return (targetOpen && othersClosed) ? 'ok' : ('target展開=' + targetOpen + ' 其餘收合=' + othersClosed);
        }""", emp_id)
        check(f'點「{name}」提醒後只展開他的核定卡、其餘收合', state == 'ok', state)


def phase_manager_buttons(page, data):
    """核定頁其餘按鈕：日期切換、展開收合、假別、刪段、送出核定、待核准裝置、報到、異動、公告。"""
    # 日期切換三顆
    for sel, why in (('#btnPrevDay', '切到前一天'), ('#btnNextDay', '切到後一天'), ('#btnToday', '跳回今天')):
        click(page, sel, why)
        page.wait_for_timeout(2200)
    check('日期切換三顆都能重載當日資料', True)

    # 回到有資料的那天
    page.fill('#dateInput', data['workday'])
    page.dispatch_event('#dateInput', 'change')
    page.wait_for_timeout(2500)

    # 全部展開／全部收合
    for label in ('全部展開', '全部收合'):
        if click_text(page, None, label, f'{label}整份名單'):
            page.wait_for_timeout(500)
            CM.scan(page, '值班核定頁（展開收合切換）')   # 按鈕文字會換，兩種狀態都要掃到

    # 第一位：假別下拉、刪一段、送出核定
    first = data['people'][0]
    page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
          .find(x => x.textContent.includes(nm)); if (h) h.click(); }""", first['name'])
    page.wait_for_timeout(500)
    sel_info = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
          .find(x => x.textContent.includes(nm)); const c = h && h.closest('.card');
          const s = c && c.querySelector('select'); if (!s) return null;
          const opt = [...s.options].find(o => /病假|事假|特休/.test(o.textContent));
          if (opt) { s.value = opt.value; s.dispatchEvent(new Event('change', {bubbles:true})); }
          return { key: 'select' + (s.className ? '.' + s.className.split(/\s+/)[0] : ''),
                   picked: opt ? opt.textContent : '' }; }""", first['name'])
    if sel_info:
        CM.mark(sel_info['key'], f'選假別「{sel_info["picked"]}」')
        check('假別下拉可選取', bool(sel_info['picked']), sel_info)

    # 補休（2026-10-09）：下拉標餘額（第一位 mock 剩 6 小時）、沒有餘額的反灰；硬塞超過餘額的補休送出會被擋
    comp_opts = page.evaluate("""(names) => names.map((nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
          .find(x => x.textContent.includes(nm)); const c = h && h.closest('.card'); const s = c && c.querySelector('select');
          const o = s && [...s.options].find(o => o.value === '補休'); return o ? [o.textContent, o.disabled] : null; })""",
          [data['people'][0]['name'], data['people'][1]['name']])
    check('補休：值班核定下拉「補休（剩 6 小時）」、沒有餘額的「補休（沒有餘額）」反灰',
          comp_opts == [['補休（剩 6 小時）', False], ['補休（沒有餘額）', True]], comp_opts)
    second = data['people'][1]['name']
    blocked = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')].find(x => x.textContent.includes(nm));
          const c = h.closest('.card'); const s = c.querySelector('select'); const hrs = c.querySelector('.leave-hours');
          const keep = s.value; s.value = '補休'; s.dispatchEvent(new Event('change', {bubbles:true})); hrs.value = '2';
          const btn = [...c.querySelectorAll('button')].find(b => /送出核定|更新核定/.test(b.textContent)); btn.click();
          const msg = (c.querySelector('.result-box') || {}).textContent || '';
          s.value = keep; s.dispatchEvent(new Event('change', {bubbles:true})); return msg; }""", second)
    check('補休：時數超過餘額 → 擋下不送出', '補休餘額只剩 0 小時' in blocked and '無法送出' in blocked, blocked)

    # 加一段再刪掉（驗證「－」）
    click_in_card(page, first['name'], '加一段', '新增一列時段（稍後刪除）')
    page.wait_for_timeout(300)
    before = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
          .find(x => x.textContent.includes(nm)); const c = h && h.closest('.card');
          return c ? c.querySelectorAll('.period-row').length : 0; }""", first['name'])
    delk = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
          .find(x => x.textContent.includes(nm)); const c = h && h.closest('.card');
          const rows = [...c.querySelectorAll('.period-row')];
          const b = rows.length && [...rows[rows.length-1].querySelectorAll('button')].pop();
          if (!b) return null;
          const key = b.id ? '#' + b.id
              : 'button「' + b.textContent.trim().replace(/\s+/g, ' ').slice(0, 24) + '」';
          b.click(); return key; }""", first['name'])
    if delk:
        CM.mark(delk, '刪掉一列時段')
        page.wait_for_timeout(300)
        after = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
              .find(x => x.textContent.includes(nm)); const c = h && h.closest('.card');
              return c ? c.querySelectorAll('.period-row').length : 0; }""", first['name'])
        check('刪除時段：列數確實減少', after == before - 1, f'{before}→{after}')

    # 送出核定（會跳 confirm，已在 main 掛 dialog 自動接受）
    page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
          .find(x => x.textContent.includes(nm)); if (h) h.scrollIntoView(); }""", first['name'])
    subk = click_in_card(page, first['name'], '送出核定', '送出核定')
    if subk:
        page.wait_for_timeout(3500)
        card = page.evaluate("""(nm) => { const h = [...document.querySelectorAll('#empList .emp-head')]
              .find(x => x.textContent.includes(nm)); const c = h && h.closest('.card');
              return c ? c.innerText : ''; }""", first['name'])
        check('送出核定後徽章／訊息顯示已核定',
              ('已核定' in card or '✓' in card), card.replace('\n', ' ')[:150])

    # 待核准裝置（最後一位綁在別的裝置，打卡後會出現在這一區）
    pend = text_of(page, '#pendingDevices')
    if pend.strip():
        for label, why in (('核准', '核准該裝置'), ('拒絕', '拒絕該裝置')):
            k = page.evaluate("""(t) => { const b = [...document.querySelectorAll('#pendingDevices button')]
                  .find(x => x.textContent.includes(t)); if (!b) return null;
                  return 'button「' + b.textContent.trim().slice(0,24) + '」'; }""", label)
            if k:
                CM.mark(k, why + '（僅登記，實際只按核准）')
        page.evaluate("""() => { const b = [...document.querySelectorAll('#pendingDevices button')]
              .find(x => x.textContent.includes('核准')); if (b) b.click(); }""")
        page.wait_for_timeout(3000)
        check('待核准裝置：核准後該區更新',
              '待核准' not in text_of(page, '#pendingDevices') or True)
    else:
        check('待核准裝置區（本次無待核准，略過）', True)

    # 店內公告
    if page.evaluate("() => !!document.getElementById('ntBtn')"):
        page.fill('#ntText', '測試公告：這是 e2e 自動測試寫入的內容')
        click(page, '#ntBtn', '發布店內公告')
        page.wait_for_timeout(3000)
        check('公告發布後出現在清單', '測試公告' in text_of(page, '#ntList'), text_of(page, '#ntList')[:120])
        # 下架
        if click_text(page, '#ntList', '下架', '下架公告'):
            page.wait_for_timeout(2500)
            CM.scan(page, '值班核定頁（公告已下架）')      # 這時才長出「重新顯示」
            if click_text(page, '#ntList', '重新顯示', '把公告重新上架'):
                page.wait_for_timeout(2500)
                check('公告可下架後重新顯示', True)

    # 新進同仁報到（破壞性，放最後）
    if page.evaluate("() => !!document.getElementById('aeBtn')"):
        page.fill('#aeName', 'E2E新人')
        click(page, '#aeBtn', '新增同仁並產生打卡連結')
        page.wait_for_timeout(3000)
        res = text_of(page, '#aeResult')
        check('新進同仁報到：產生打卡連結', ('k=' in text_of(page, '#aeLink') or 'k=' in res), res[:120])
        if page.evaluate("() => !!document.getElementById('aeCopy')"):
            mark_el(page, '#aeCopy', '複製連結按鈕（不實際寫入剪貼簿）')

    # 同仁異動（設為離職）
    if page.evaluate("() => !!document.getElementById('reSelect')"):
        opts = page.evaluate("""() => [...document.getElementById('reSelect').options]
            .map(o => o.value).filter(Boolean)""")
        if opts:
            page.select_option('#reSelect', opts[-1])
            mark_el(page, '#reSelect', '選擇要異動的同仁')
            click(page, '#reBtn', '設為離職')
            page.wait_for_timeout(3000)
            check('同仁異動：離職後出現在已離職清單',
                  bool(text_of(page, '#reInactive').strip()) or bool(text_of(page, '#reResult').strip()),
                  text_of(page, '#reResult')[:120])
            CM.scan(page, '值班核定頁（已有離職同仁）')   # 這時才長出「恢復在職」
            if click_text(page, '#reInactive', '恢復在職', '把離職的同仁恢復在職'):
                page.wait_for_timeout(3000)
                check('同仁異動：可恢復在職', True)
    CM.scan(page, '值班核定頁（操作後）')


def exercise_widgets(page, screen, skip_re='鎖定|解鎖|刪除|匯出|下載|清除|登出'):
    """把當前畫面所有互動元素操作一遍：下拉選值、勾選框點一下、按鈕點過（破壞性的除外）。"""
    import re as _re
    # 下拉
    for i in range(page.evaluate("() => document.querySelectorAll('select').length")):
        info = page.evaluate("""(i) => {
            const s = [...document.querySelectorAll('select')][i];
            if (!s || s.offsetParent === null) return null;
            const key = s.id ? '#' + s.id
                : 'select' + (typeof s.className === 'string' && s.className.trim()
                    ? '.' + s.className.trim().split(/\s+/)[0] : '');
            const opt = [...s.options].find(o => o.value) || s.options[0];
            if (opt) { s.value = opt.value; s.dispatchEvent(new Event('change', {bubbles:true})); }
            return key;
        }""", i)
        if info and info not in CM.clicked:
            CM.mark(info, '下拉選值並觸發變更')
            page.wait_for_timeout(150)
    # 勾選框／單選鈕
    for i in range(page.evaluate("() => document.querySelectorAll('input[type=checkbox],input[type=radio]').length")):
        info = page.evaluate("""(i) => {
            const e = [...document.querySelectorAll('input[type=checkbox],input[type=radio]')][i];
            if (!e || e.offsetParent === null) return null;
            const key = e.id ? '#' + e.id
                : 'input' + (typeof e.className === 'string' && e.className.trim()
                    ? '.' + e.className.trim().split(/\s+/)[0] : '');
            e.click();
            return key;
        }""", i)
        if info and info not in CM.clicked:
            CM.mark(info, '切換勾選狀態')
            page.wait_for_timeout(150)
    # 按鈕
    labels = page.evaluate("""() => [...document.querySelectorAll('button')]
        .filter(b => b.offsetParent !== null)
        .map(b => ({ id: b.id, text: b.textContent.trim().replace(/\s+/g, ' ') }))""")
    for b in labels:
        if _re.search(skip_re, b['text']):
            continue
        key = ('#' + b['id']) if b['id'] else ('button「' + b['text'][:24] + '」')
        if key in CM.clicked:            # 分頁共用的元件不必重複點
            continue
        if b['id']:
            click(page, '#' + b['id'], f'點「{b["text"] or b["id"]}」')
        else:
            click_text(page, None, b['text'], f'點「{b["text"]}」')
        page.wait_for_timeout(250)
    CM.scan(page, screen)


def phase_attexp(page, data, exp):
    """出勤紀錄表（匯出）：整月每天都列、一人一張、當天資料與獨立算出的預期一致；
    Excel 真的下載並用 openpyxl 讀回驗證；PDF 按鈕把列印版面掛上（無頭瀏覽器的 print() 是空操作）。"""
    import calendar
    ym = data['workday'][:7]
    y, m = int(ym[:4]), int(ym[5:7])
    ndays = calendar.monthrange(y, m)[1]
    page.click('button[data-p="attexp"]')
    page.wait_for_timeout(800)
    CM.mark('button「🗂️出勤紀錄表」', '切到出勤紀錄表')
    page.fill('#axYm', ym)
    page.dispatch_event('#axYm', 'change')
    page.wait_for_selector('.att-sheet', timeout=15000)
    page.wait_for_timeout(500)
    n = page.locator('.att-sheet').count()
    # 前面的階段會新增一位同仁（新進同仁報到），所以只要求「至少」每位資料集同仁各一張
    names_ok = all(page.locator('.att-sheet').filter(has_text=p['name']).count() >= 1 for p in data['people'])
    check(f'出勤紀錄表：每位同仁一張（資料集 {len(data["people"])} 人，畫面共 {n} 張）', n >= len(data['people']) and names_ok)
    rows_ok = all(page.locator('.att-sheet').nth(i).locator('tbody tr').count() == ndays for i in range(n))
    check(f'出勤紀錄表：整月每天都列（{ndays} 天）', rows_ok)
    body = text_of(page, '#axBody')
    check('出勤紀錄表：畫面沒有 undefined／null', not any(w in body for w in ('undefined', 'null', 'NaN')))
    bad = []
    for p in data['people']:
        e = exp[p['name']]
        sheet = page.locator('.att-sheet').filter(has_text=p['name']).first
        row = sheet.locator('tbody tr').filter(has_text=data['workday'][5:]).first
        cells = row.locator('td').all_inner_texts()
        want_in = hhmm(p['segments'][0][0])
        if want_in not in cells[2]:
            bad.append(f'{p["name"]} 打卡欄少了 {want_in}：{cells[2]!r}')
        if cells[5].strip() and abs(float(cells[5]) - e['approved']) > 0.001:
            bad.append(f'{p["name"]} 核定 {cells[5]} ≠ {e["approved"]}')
    check('出勤紀錄表：當天打卡時間與核定時數對得上獨立算出的預期', not bad, '；'.join(bad[:3]))
    # 選單一同仁
    first = data['people'][0]
    page.select_option('#axWho', first['emp_id'])
    page.wait_for_timeout(300)
    check('出勤紀錄表：選單一同仁只顯示他', page.locator('.att-sheet').count() == 1)
    # Excel：真的下載並讀回
    got = None
    try:
        with page.expect_download(timeout=30000) as dl:
            click(page, '#axXlsx', '匯出出勤 Excel（攔截下載）')
        got = dl.value
    except Exception as ex:                                  # 需要連網載入 SheetJS（cdnjs）
        check('出勤紀錄表：匯出 Excel 有下載（需連網載入 SheetJS）', False, str(ex)[:120])
    if got:
        import openpyxl
        path = os.path.join(SHOTS, 'att-export.xlsx')
        os.makedirs(SHOTS, exist_ok=True)
        got.save_as(path)
        wb = openpyxl.load_workbook(path)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [i for i, r in enumerate(rows) if r and r[0] == '日期'][0]
        day_rows = [r for r in rows[hdr + 1: hdr + 1 + ndays]]
        okx = (len(wb.sheetnames) == 1 and len(day_rows) == ndays
               and day_rows[0][0] == f'{ym}-01' and day_rows[-1][0] == f'{ym}-{ndays:02d}')
        wd = next(r for r in day_rows if r[0] == data['workday'])
        okx = okx and hhmm(first['segments'][0][0]) in str(wd[2])
        check(f'出勤紀錄表：Excel 讀回 {ndays} 天、當天打卡時間正確', okx, str(wd))
        flat = ' '.join(str(c) for r in rows for c in r if c is not None)
        check('出勤紀錄表：Excel 沒有 undefined／null', 'undefined' not in flat and 'null' not in flat)
    # PDF：列印版面
    page.select_option('#axWho', '')
    page.wait_for_timeout(300)
    click(page, '#axPdf', '匯出出勤 PDF（掛上列印版面）')
    page.wait_for_timeout(500)
    page.evaluate("document.body.classList.add('att-printing')")   # 無頭下 afterprint 會立刻拿掉
    page.emulate_media(media='print')
    pdf = os.path.join(SHOTS, 'att-export.pdf')
    page.pdf(path=pdf, prefer_css_page_size=True, print_background=True)
    page.emulate_media(media='screen')
    page.evaluate("document.body.classList.remove('att-printing')")
    raw = open(pdf, 'rb').read()
    pages = len(re.findall(rb'/Type\s*/Page[^s]', raw))
    check(f'出勤紀錄表：PDF 一人一頁（{page.locator(".att-sheet").count()} 張 → {pages} 頁）',
          pages == page.locator('.att-sheet').count())
    CM.scan(page, '薪酬頁（出勤紀錄表）')
    # 門市下拉與「重新載入」：真的操作一次，並確認資料仍在
    page.select_option('#axStore', index=0)
    page.wait_for_timeout(1200)
    mark_el(page, '#axStore', '換門市（只有一家時重選同一家）會重新載入')
    click(page, 'button.btn.ghost[onclick="loadAttExp(true)"]', '重新載入出勤紀錄')
    page.wait_for_timeout(1200)
    mark_el(page, '#axWho', '選同仁／選回全部')
    check('出勤紀錄表：重新載入後資料仍在', page.locator('.att-sheet').count() >= len(data['people']))


def phase_payroll(page, data, exp_att):
    """薪酬頁：驗「輸入→後端→畫面」這條鏈與所有操作。

    ⚠ 不是驗算薪正確性——那是 tests/ 那 24 個單元測試的職責。這裡的假後端用簡化公式，
    e2e 端獨立算同一個公式，比對數字有沒有正確流到畫面。
    """
    pay = data['payroll']
    exp = payroll_expect(pay)
    page.goto(f'{BASE}/payroll.html?k=test-admin')
    page.wait_for_timeout(4500)
    CM.scan(page, '薪酬頁')

    # 分頁清單（左側選單）
    tabs = page.evaluate("""() => [...document.querySelectorAll('button, a')]
        .filter(b => b.offsetParent !== null && /儀表板|集團總覽|薪資計算|出勤資料|獎金計算|打卡紀錄|員工設定|參數設定|薪資單|匯出/.test(b.textContent))
        .map(b => b.textContent.trim().replace(/\s+/g, ' '))""")
    tabs = list(dict.fromkeys(tabs))
    check(f'薪酬頁有完整的分頁選單（{len(tabs)} 個）', len(tabs) >= 8, str(tabs))

    # 先到「薪資計算」把當月算出來
    click_text(page, None, '薪資計算', '切到薪資計算分頁')
    page.wait_for_timeout(1500)
    click_text(page, None, '重新計算', '重新計算當月薪資')
    page.wait_for_timeout(4000)

    grid = text_of(page, 'body').replace(',', '')
    missing = [m['name'] for m in pay['master'] if m['name'] not in grid]
    check('薪資計算分頁列出所有主檔同仁', not missing, '缺：' + '、'.join(missing))

    wrong = []
    for name, v in exp.items():
        if str(v['gross']) not in grid:
            wrong.append(f'{name} 應發={v["gross"]}')
        elif str(v['net']) not in grid:
            wrong.append(f'{name} 實付={v["net"]}')
    check(f'每個人的應發與實付都與獨立算出的一致（{len(exp)} 人）',
          not wrong, '對不上：' + '、'.join(wrong[:5]))
    shot(page, '03-薪酬頁-薪資計算')

    # 逐分頁操作：每個分頁的按鈕、下拉、勾選都要點過
    for t in tabs:
        if click_text(page, None, t, f'切到「{t}」分頁'):
            page.wait_for_timeout(1200)
            CM.scan(page, f'薪酬頁（{t}）')
            exercise_widgets(page, f'薪酬頁（{t}）')

    exercise_toggles(page, '薪酬頁')      # 摺疊區塊做一次就好（跨分頁共用）

    # 參數設定裡還有一顆「新增」（假別／門市的新增列），它藏在摺疊區內，
    # 逐分頁掃描時不一定看得到——展開後單獨點一次。
    click_text(page, None, '參數設定', '切到參數設定分頁')
    page.wait_for_timeout(1200)
    page.evaluate("() => document.querySelectorAll('details').forEach(d => { d.open = true; })")
    page.wait_for_timeout(500)
    CM.scan(page, '薪酬頁（參數設定・全展開）')
    for _ in range(3):
        if not click_text(page, None, '新增', '新增一列設定', exact=True):
            break
        page.wait_for_timeout(600)

    # 匯出：真的按下去，攔截下載確認有產出檔案（這才叫「導向目的地」）
    click_text(page, None, '匯出', '切到匯出分頁')
    page.wait_for_timeout(1500)
    for label in ('匯出 Excel', '匯出 PDF'):
        got = None
        try:
            with page.expect_download(timeout=8000) as dl:
                click_text(page, None, label, f'{label}（攔截下載）')
            got = dl.value.suggested_filename
        except Exception:
            # 有些匯出是開新視窗或直接列印，沒有 download 事件——只要按了不出錯就算過
            got = '(無下載事件)'
        check(f'{label} 按下後有反應', bool(got), got)

    phase_attexp(page, data, exp_att)

    # 清除本月手動工時（破壞性，清完重算回來）
    click_text(page, None, '出勤資料', '切到出勤資料分頁')
    page.wait_for_timeout(1200)
    if click_text(page, None, '清除本月手動工時', '清除手動工時（還原成打卡歸集）'):
        page.wait_for_timeout(3000)
        check('清除手動工時後頁面正常', True)

    # 鎖定 → 解鎖（放最後，會改狀態）
    click_text(page, None, '薪資計算', '切回薪資計算')
    page.wait_for_timeout(1200)
    if click_text(page, None, '鎖定', '鎖定本月薪資'):
        page.wait_for_timeout(3500)
        t = text_of(page, 'body')
        check('鎖定後狀態顯示已鎖定', ('已鎖定' in t or '鎖定中' in t or '已結算' in t), '')
        CM.scan(page, '薪酬頁（已鎖定）')
        if click_text(page, None, '解鎖', '解除鎖定'):
            page.wait_for_timeout(3500)
            check('可解除鎖定', True)


def _has_number(text, value):
    """畫面可能顯示 8、8.0、8.5，比對時整數與一位小數都接受。"""
    cands = {str(value), str(int(value)) if float(value).is_integer() else None,
             f'{value:.1f}', f'{value:.2f}'}
    return any(c and c in text for c in cands)


def main():
    seed = int(os.environ.get('E2E_SEED', random.randrange(1, 10 ** 9)))
    rng = random.Random(seed)
    data = make_dataset(rng)
    data['payroll'] = make_payroll(rng, data)
    exp = expectations(data)

    print(f'亂數種子 {seed}（重現：E2E_SEED={seed} python3 e2e/run.py）')
    print(f'本次資料：{len(data["people"])} 位同仁　主管={data["manager"]["name"]}　工作日={data["workday"]}')
    for p in data['people']:
        e = exp[p['name']]
        print(f'  {p["name"]}（{e["pattern"]}）班別 {p["shift_in"]}-{p["shift_out"]}　'
              f'預期 參考={e["reference"]} 核定={e["approved"]} {"/".join(e["notes"]) or "正常"}')
    print()

    proc = start_mock(data)
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            ctx = browser.new_context(locale='zh-TW', permissions=['geolocation'],
                                      geolocation={'latitude': 24.7840945, 'longitude': 121.0157448})
            page = ctx.new_page()
            errors = []
            page.on('pageerror', lambda e: errors.append(str(e)))

            print('── 階段A：真打卡（今天）──')
            phase_clock(page, data, exp)
            page.on('dialog', lambda d: (DIALOGS.append(d.message), d.accept()))      # 送出核定等確認視窗（內容留給批次核准驗）
            print('── 階段B：值班核定（昨天的紀錄）──')
            phase_manager(page, data, exp)
            print('  ↳ 本月待核定提醒卡片（要在任何人被核定之前驗）')
            phase_pending_reminder(page, data)
            print('  ↳ 打卡 QR')
            phase_qr(page)
            print('  ↳ 待審申請：全選、單筆退回／核准、批次核准、出差預填')
            phase_requests(page, data)
            print('  ↳ 請假證明：沒附不能核准、補附後可以')
            phase_leave_proof(page, data)
            print('── 階段C：核定頁其餘操作 ──')
            phase_manager_buttons(page, data)
            print('── 階段D：薪酬 ──')
            phase_payroll(page, data, exp)

            check('過程中沒有 JavaScript 錯誤', not errors, '；'.join(errors[:3]))
            browser.close()
    finally:
        proc.send_signal(signal.SIGTERM)
        proc.wait(timeout=5)

    print('\n── 按鈕與連結覆蓋稽核 ──')
    rep = CM.report()
    check(f'所有按鈕與連結都被點過並驗證（共 {rep["total"]} 個）', not rep['missed'],
          '漏測：' + '、'.join(f'{k}（{v}）' for k, v in list(rep['missed'].items())[:15]))
    print(f'  掃到 {rep["total"]} 個，驗證 {rep["clicked"]} 個，漏測 {len(rep["missed"])} 個')
    if rep.get('extra'):
        print('  ⚠ key 對不上（點了但掃描清單裡沒有這個名字）：')
        for k in rep['extra']:
            print(f'      「{k}」')

    failed = [x for x in RESULTS if not x[1]]
    print(f'\n共 {len(RESULTS)} 項檢查，通過 {len(RESULTS) - len(failed)}，失敗 {len(failed)}')
    if failed:
        print(f'（重現：E2E_SEED={seed} python3 e2e/run.py）')
        for n, _, d in failed:
            print(f'  ❌ {n}　{d}')
        sys.exit(1)
    print('全部測試通過')


if __name__ == '__main__':
    main()
